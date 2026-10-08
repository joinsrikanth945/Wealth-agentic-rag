from docx import Document as DocxDocument

from app.services.ingestion import chunk_documents, load_file

SAMPLE_TEXT = "Advisors can view client portfolios in the Front Office module."


def write_text_file(tmp_path, name, text):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


# ---------- Loading ----------

def test_load_txt_file_returns_its_text(tmp_path):
    path = write_text_file(tmp_path, "guide.txt", SAMPLE_TEXT)
    docs = load_file(path)
    assert len(docs) >= 1
    assert "Front Office" in docs[0].page_content


def test_load_markdown_file_returns_its_text(tmp_path):
    path = write_text_file(tmp_path, "guide.md", "# MFA setup\n\nScan the QR code with the app.")
    docs = load_file(path)
    assert "QR code" in " ".join(d.page_content for d in docs)


def test_load_docx_file_returns_its_text(tmp_path):
    path = tmp_path / "guide.docx"
    doc = DocxDocument()
    doc.add_paragraph("Reset a user token from the admin console.")
    doc.save(path)
    docs = load_file(path)
    assert "admin console" in " ".join(d.page_content for d in docs)


def test_loaded_document_records_its_source(tmp_path):
    path = write_text_file(tmp_path, "guide.txt", SAMPLE_TEXT)
    docs = load_file(path)
    assert "guide.txt" in str(docs[0].metadata.get("source", ""))


# ---------- Chunking ----------

def test_long_document_is_split_into_several_chunks(tmp_path):
    long_text = "\n\n".join(f"Paragraph {i}: {SAMPLE_TEXT}" for i in range(200))
    path = write_text_file(tmp_path, "long.txt", long_text)
    chunks = chunk_documents(load_file(path))
    assert len(chunks) > 1


def test_chunks_are_never_empty(tmp_path):
    long_text = "\n\n".join(f"Paragraph {i}: {SAMPLE_TEXT}" for i in range(200))
    path = write_text_file(tmp_path, "long.txt", long_text)
    chunks = chunk_documents(load_file(path))
    assert all(c.page_content.strip() for c in chunks)


def test_chunks_keep_the_source_of_their_document(tmp_path):
    long_text = "\n\n".join(f"Paragraph {i}: {SAMPLE_TEXT}" for i in range(200))
    path = write_text_file(tmp_path, "long.txt", long_text)
    chunks = chunk_documents(load_file(path))
    assert all("long.txt" in str(c.metadata.get("source", "")) for c in chunks)


def test_no_documents_gives_no_chunks():
    assert chunk_documents([]) == []


# ---------- OCR (scanned documents and images) ----------

import shutil

import pytest
from langchain_core.documents import Document
from PIL import Image, ImageDraw, ImageFont

from app.services import ingestion

requires_tesseract = pytest.mark.skipif(
    shutil.which("tesseract") is None, reason="Tesseract OCR is not installed"
)

SCANNED_TEXT = "Client activation links are valid for 72 hours"


def make_scanned_page(text=SCANNED_TEXT):
    """Create an image of a printed page: black text on white, like a scan."""
    image = Image.new("RGB", (1700, 300), "white")
    font = ImageFont.load_default(size=48)
    ImageDraw.Draw(image).text((60, 110), text, fill="black", font=font)
    return image


@requires_tesseract
def test_image_file_is_read_with_ocr(tmp_path):
    path = tmp_path / "scan.png"
    make_scanned_page().save(path)

    docs = load_file(path)

    assert "72 hours" in docs[0].page_content
    assert docs[0].metadata["ocr"] is True


@requires_tesseract
def test_scanned_pdf_page_is_read_with_ocr(tmp_path):
    path = tmp_path / "scanned.pdf"
    make_scanned_page().save(path, "PDF")   # an image-only PDF, with no real text

    docs = load_file(path)

    assert "activation" in docs[0].page_content.lower()
    assert docs[0].metadata["ocr"] is True


def test_pdf_with_real_text_skips_ocr(mocker, tmp_path):
    text_page = Document(page_content="This page already has plenty of real text.", metadata={"page": 0})
    mocker.patch.object(ingestion, "PyPDFLoader").return_value.load.return_value = [text_page]
    ocr = mocker.patch.object(ingestion, "ocr_image")

    docs = load_file(tmp_path / "normal.pdf")

    ocr.assert_not_called()                  # the slow OCR path is never used
    assert "ocr" not in docs[0].metadata


def test_image_files_are_supported_for_upload():
    assert {".png", ".jpg", ".jpeg"} <= ingestion.SUPPORTED
