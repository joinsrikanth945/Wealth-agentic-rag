from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PIL import Image

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
SUPPORTED = {".pdf", ".txt", ".md", ".docx"} | IMAGE_SUFFIXES

# A PDF page with less extractable text than this is treated as scanned and OCR'd
OCR_MIN_CHARS = 20
# Resolution used to render scanned PDF pages for OCR (300 DPI is standard for OCR)
OCR_DPI = 300


def ocr_image(image: Image.Image) -> str:
    """Read the text in an image with Tesseract OCR."""
    import pytesseract  # imported here so the app still starts if OCR is never used

    return pytesseract.image_to_string(image).strip()


def _load_pdf(path: Path) -> list[Document]:
    """Extract text from a PDF; pages with no real text (scanned pages) are OCR'd."""
    docs = PyPDFLoader(str(path)).load()
    scanned = [d for d in docs if len(d.page_content.strip()) < OCR_MIN_CHARS]
    if not scanned:
        return docs  # normal PDF: no OCR needed, fast path

    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        for i, doc in enumerate(docs):
            if doc in scanned:
                page_number = doc.metadata.get("page", i)
                image = pdf[page_number].render(scale=OCR_DPI / 72).to_pil()
                doc.page_content = ocr_image(image)
                doc.metadata["ocr"] = True
    finally:
        pdf.close()
    return docs


def _load_image(path: Path) -> list[Document]:
    """OCR an image file, such as a photo or scan of a printed page."""
    with Image.open(path) as image:
        text = ocr_image(image)
    return [Document(page_content=text, metadata={"source": str(path), "ocr": True})]


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix in IMAGE_SUFFIXES:
        return _load_image(path)
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), encoding="utf-8").load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))
